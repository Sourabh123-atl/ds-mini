#!/usr/bin/env python3
"""Build the Data Science Mini Project report (SLRTCE template).
  .venv/bin/python scripts-report/build_report.py
  libreoffice --headless --convert-to pdf --outdir report report/DS_Mini_Project_Report.docx
"""
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).parent))
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, Cm
from PIL import Image
from report_helpers import add_para, add_mixed_para, add_code_block, section_heading

ROOT = Path(__file__).resolve().parents[1]
FIGS = ROOT / "report" / "figures"
OUT = ROOT / "report" / "DS_Mini_Project_Report.docx"
SUBJECT = "Data Science"
TITLE = "COVID/Influenza Trend Analysis Using Social Media Data"
GUIDE = "Asst. Prof. Sheetal Dalvi"          # change here if the guide changes
YEAR = "2026-27"
BLANK = "(" + "\u00a0" * 12 + ")"                          # roll number to be written by hand
NAMES = [f"Sourabh Patel\u00a0{BLANK}", f"Chirag Khaire\u00a0{BLANK}", f"Khubaib Khan\u00a0{BLANK}"]


def body(doc, t):
    add_para(doc, t, size=12, align="justify", space_after=8, line_spacing=1.4)


def sub(doc, t):
    add_para(doc, t, size=12, bold=True, align="left", space_before=6, space_after=4)


def bullet(doc, name, text):
    add_mixed_para(doc, [(f"\u2022  {name} \u2014 ", True), (text, False)], align="justify", size=12, space_after=6, line_spacing=1.35)


def dot(doc, t):
    add_para(doc, "\u2022  " + t, size=12, align="justify", space_after=3, line_spacing=1.3)


def figure(doc, f, cap):
    w, h = Image.open(FIGS / f).size
    width = min(5.8, 8.3 * w / h)          # full width, but never taller than the page
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(FIGS / f), width=Inches(width))
    cp = doc.add_paragraph(); cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_after = Pt(6)
    r = cp.add_run(cap); r.italic = True; r.font.size = Pt(10); r.font.name = "Liberation Serif"


def build():
    doc = Document()
    for s in doc.sections:
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Inches(1.0)

    # COVER
    add_para(doc, "", space_after=12)
    add_para(doc, f"{SUBJECT} Mini Project Report", size=16, bold=True, align="center", space_after=16, line_spacing=1.0)
    add_para(doc, TITLE, size=20, bold=True, align="center", space_before=4, space_after=16)
    add_para(doc, "Submitted fulfilment of the requirements of the", size=12, align="center", space_after=0)
    add_para(doc, f"{SUBJECT} Mini Project Semester VII in degree of", size=12, align="center", space_after=0)
    add_para(doc, "Bachelor of Computer Engineering", size=12, align="center", space_after=12)
    add_para(doc, "BY", size=12, align="center", space_after=6, line_spacing=1.0)
    for n in NAMES:
        add_para(doc, n, size=13, bold=True, align="center", space_after=2)
    add_para(doc, "", space_after=8)
    add_para(doc, "Guide", size=12, align="center", space_after=2, line_spacing=1.0)
    add_para(doc, GUIDE, size=12, align="center", space_after=12)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(8)
    p.add_run().add_picture(str(FIGS / "college-logo.png"), width=Inches(1.55))
    add_para(doc, "Department Of Computer Engineering", size=13, bold=True, align="center", space_after=2)
    add_para(doc, "Shree L. R. Tiwari College Of Engineering", size=13, bold=True, align="center", space_after=2)
    add_para(doc, "Kanakia Park, Mira Road (E), Thane -401 107, Maharashtra.", size=12, bold=True, align="center", space_after=12)
    add_para(doc, YEAR, size=14, bold=True, align="center", space_after=0, line_spacing=1.0)

    # APPROVAL
    add_para(doc, f"{SUBJECT} Mini Project Report Approval", size=16, bold=True, align="center", space_after=28, page_break_before=True)
    add_mixed_para(doc, [
        ("This project report entitled ", False), (f"\u201c{TITLE}\u201d", True), (" by ", False),
        (f"Mr. {NAMES[0]}, Mr. {NAMES[1]} and Mr. {NAMES[2]}", True),
        (f" is approved for the Mini Project Semester VII project work of B.E CMPN at Shree L.R.Tiwari College of Engineering, in the academic year 2026-2027.", False)],
        align="justify", size=12, space_after=6, line_spacing=1.5)
    for _ in range(10):
        add_para(doc, "", space_after=12)
    add_para(doc, "Internal Examiner", size=12, bold=True, align="right", space_after=0)

    # DECLARATION
    add_para(doc, "Declaration", size=16, bold=True, align="center", space_after=20, page_break_before=True)
    add_para(doc, (
        "I declare that this written submission represents my ideas in my own words and where "
        "others' ideas or words have been included, I have adequately cited and referenced the "
        "original sources. I also declare that I have adhered to all principles of academic "
        "honesty and integrity and have not misrepresented or fabricated or falsified any "
        "idea/data/fact/source in my submission. I understand that any violation of the above "
        "will be cause for disciplinary action by the Institute and can also evoke penal action "
        "from the sources which have thus not been properly cited or from whom proper permission "
        "has not been taken when needed."), size=12, align="justify", space_after=36, line_spacing=1.5)
    for n in NAMES:
        add_para(doc, "________________________", size=12, align="right", space_before=18, space_after=0, line_spacing=1.0)
        add_para(doc, "Signature", size=11, align="right", space_after=0, line_spacing=1.0)
        add_para(doc, "Mr. " + n, size=12, bold=True, align="right", space_after=6)
    add_para(doc, "", space_after=36)
    add_para(doc, "Date:", size=12, space_after=6)
    add_para(doc, "Place:", size=12, space_after=0)

    # TOC
    add_para(doc, "Table of Contents", size=14, bold=True, space_after=14, page_break_before=True)
    for it in ["1. Introduction", "2. System Architecture", "3. Implementation", "Technologies Used",
               "4. Results", "5. Conclusion", "6. References"]:
        add_para(doc, it, size=12, space_after=4, line_spacing=1.5)

    # 1 INTRODUCTION
    section_heading(doc, "1. Introduction")
    body(doc, "During the COVID-19 pandemic, social media became one of the fastest public signals of what people were "
         "worried about: case counts, deaths, testing, masks, lockdowns, vaccines and, for respiratory illness in "
         "general, symptoms such as fever and cough or mentions of the flu. Official case data is published with a "
         "delay and says nothing about public mood. Tweets are noisy, but they are timestamped, large in volume and "
         "partly geo-labelled, so they can be analysed as a time series alongside official statistics.")
    body(doc, "This Data Science mini-project analyses a well-known public Twitter dataset, the COVID19 Tweets dataset "
         "by Gabriel Preda (Kaggle), which contains 179,108 English tweets with the #covid19 hashtag posted between "
         "24 July and 30 August 2020. Each tweet is flagged as India, Maharashtra or Mumbai from the free-text "
         "user_location field, and the tweet series is compared with official daily cases for India, Maharashtra and "
         "Mumbai (covid19india.org archive) and the world (Our World in Data). A sentiment classifier is trained on a "
         "second, hand-labelled COVID tweet dataset and exported. All results are shown in an interactive Streamlit "
         "dashboard with filters and several analysis tabs.")
    sub(doc, "1.1 Objectives")
    for o in ["Acquire real, citable COVID-19 social-media data and official case time series for India, Maharashtra and Mumbai.",
              "Perform multiple trend analyses: tweet volume, sentiment, hashtags/keywords, topics, symptom/health terms, region breakdown.",
              "Measure the (lag) correlation between tweet volume and new cases.",
              "Train, evaluate and export a sentiment classification model.",
              "Present everything in an interactive dashboard with date, region, keyword and sentiment filters."]:
        dot(doc, o)
    sub(doc, "1.2 Datasets")
    bullet(doc, "COVID19 Tweets (G. Preda)", "179,108 rows, 13 columns (user_location, date, text, hashtags, source, ...). Downloaded from the author\u2019s GitHub mirror of the Kaggle dataset.")
    bullet(doc, "Coronavirus Tweets NLP (Kaggle, datatattle)", "44,955 manually tagged tweets (41,157 train + 3,798 test) with five sentiment labels, from the Hugging Face mirror argilla/twitter-coronavirus. Used only to train the model.")
    bullet(doc, "covid19india.org archive", "case_time_series.csv (641 rows, India daily), states.csv (21,675 rows), districts.csv (358,633 rows, includes Mumbai).")
    bullet(doc, "Our World in Data", "cases_deaths/full_data.csv; the World series (1,674 rows) gives global daily cases.")

    # 2 ARCHITECTURE
    section_heading(doc, "2. System Architecture")
    body(doc, "The system is an offline batch pipeline followed by a lightweight dashboard. Heavy work (sentiment scoring, "
         "topic modelling, case joins, correlation) runs once in Python and writes compact Parquet/CSV/JSON files, so "
         "the dashboard loads in seconds.")
    add_code_block(doc, """+-------------------------+   +--------------------------------+
| COVID19 Tweets (Preda)  |   | covid19india.org archive + OWID|
| 179,108 tweets (CSV)    |   | India / MH / Mumbai / World    |
+-----------+-------------+   +---------------+----------------+
            |                                 |
            v                                 v
+-----------------------------------------------------------+
| scripts/analyze.py                                        |
|  clean + dedupe -> region flags (regex on user_location)  |
|  VADER sentiment -> hashtags -> symptom/health terms      |
|  TF-IDF + NMF topics (k=8) -> daily join with cases       |
|  lag correlation (-7..+7 days) -> summary.json            |
+---------------------------+-------------------------------+
                            |   data/processed/*.parquet,csv,json
+---------------------------v-------------------------------+
| scripts/train_model.py  (Coronavirus Tweets NLP, labelled)|
|  TF-IDF(1-2 grams) + Logistic Regression -> joblib+metrics|
+---------------------------+-------------------------------+
                            v
+-----------------------------------------------------------+
| app/streamlit_app.py  (Plotly)                            |
|  Sidebar: date range, region, keyword/#hashtag, sentiment |
|  Tabs: Trends | Sentiment | Topics | Keywords | Cases     |
|        correlation | Try the model | Data                 |
+-----------------------------------------------------------+""", size=8)
    sub(doc, "2.1 Data preparation")
    body(doc, "Tweets without a valid date or text are dropped and exact duplicates (same user and text) are removed, "
         "leaving 178,484 tweets on 25 distinct days (25 July \u2013 30 August 2020; the dataset has no tweets on some "
         "days, e.g. 3, 5, 15, 19\u201321 and 23\u201328 August). Region flags are set by regular expressions on the "
         "lower-cased user_location: Mumbai (mumbai, bombay, navi mumbai, thane, andheri, bandra) = 2,488 tweets; "
         "Maharashtra (adds pune, nagpur, nashik and other cities) = 3,080; India (adds India, Indian states and "
         "major cities) = 20,582 (11.5%).")
    body(doc, "Official data: daily new cases for Maharashtra and Mumbai are the first difference of the cumulative "
         "Confirmed column. The OWID World new_cases column is reported in weekly batches, so the 7-day sum "
         "divided by 7 is used as the global daily value.")
    sub(doc, "2.2 Analyses performed")
    for n, r in [
        ("Volume trend", "daily and weekly tweet counts, overall and per region, plotted against new cases."),
        ("Sentiment trend", "VADER compound score per tweet (Positive \u2265 0.05, Negative \u2264 \u22120.05), daily mix and regional comparison."),
        ("Hashtags and keywords", "top hashtags overall and by region; free-text and #hashtag search."),
        ("Symptom / health terms", "fever, cough, vaccine, oxygen, lockdown, flu/influenza, mask, hospital, test, death, as % of tweets per day and India vs rest of world."),
        ("Topic modelling", "TF-IDF (6,000 terms) + Non-negative Matrix Factorisation with 8 topics; topic share by week and month."),
        ("Lag correlation", "Pearson r between tweets on day d and new cases on day d + lag, for lags \u22127 to +7."),
        ("Sentiment model", "TF-IDF + logistic regression classifier trained on a labelled COVID tweet corpus and applied to all tweets."),
    ]:
        bullet(doc, n, r)

    # 3 IMPLEMENTATION
    section_heading(doc, "3. Implementation")
    sub(doc, "3.1 Region flags, sentiment and terms (scripts/analyze.py)")
    add_code_block(doc, r'''MUMBAI = r"\b(mumbai|bombay|navi mumbai|thane|andheri|bandra)\b"
MAHA   = r"\b(maharashtra|pune|nagpur|nashik|aurangabad|...)\b|" + MUMBAI
INDIA  = r"\b(india|bharat|delhi|bengaluru|chennai|kolkata|...)\b|" + MAHA
loc = t["user_location"].fillna("").str.lower()
t["is_mumbai"] = loc.str.contains(MUMBAI); t["is_india"] = loc.str.contains(INDIA)

sia = SentimentIntensityAnalyzer()                         # VADER
t["compound"]  = [sia.polarity_scores(x)["compound"] for x in t["text"]]
t["sentiment"] = pd.cut(t["compound"], [-1.01, -0.05, 0.05, 1.01],
                        labels=["Negative", "Neutral", "Positive"])
TERMS = {"fever": r"\bfever", "vaccine": r"\bvaccin", "oxygen": r"\boxygen",
         "lockdown": r"\block ?down", "flu/influenza": r"\b(flu|influenza)\b", ...}
for k, rx in TERMS.items():
    t["t_" + k] = t["text"].str.lower().str.contains(rx)''', size=8)
    sub(doc, "3.2 Topic modelling with NMF")
    add_code_block(doc, r'''vec = TfidfVectorizer(max_features=6000, min_df=20, max_df=0.4,
                      stop_words=STOP_WORDS + ["covid", "coronavirus", ...])
X   = vec.fit_transform(clean_text)
nmf = NMF(n_components=8, random_state=42, init="nndsvda", max_iter=300)
W   = nmf.fit_transform(X)                 # tweet x topic weights
t["topic"] = W.argmax(axis=1)              # dominant topic per tweet''', size=8)
    sub(doc, "3.3 Joining official cases and lag correlation")
    add_code_block(doc, r'''mh = st[st.State == "Maharashtra"].sort_values("Date")
mh = pd.DataFrame({"day": mh.Date, "maha_cases": mh.Confirmed.diff()})
mu = ds[(ds.State == "Maharashtra") & (ds.District == "Mumbai")]
ow["world_cases_owid"] = ow["weekly_cases"] / 7      # OWID reports weekly
for L in range(-7, 8):                    # tweets on day d vs cases on d+L
    y = cases[cname].reindex(daily["day"] + pd.Timedelta(days=L)).values
    r = np.corrcoef(daily[col][~np.isnan(y)], y[~np.isnan(y)])[0, 1]''', size=8)
    sub(doc, "3.4 Sentiment model training and export (scripts/train_model.py)")
    add_code_block(doc, r'''MAP = {"Extremely Negative": "Negative", "Negative": "Negative",
       "Neutral": "Neutral", "Positive": "Positive", "Extremely Positive": "Positive"}
pipe = Pipeline([
  ("tfidf", TfidfVectorizer(ngram_range=(1, 2), min_df=2,
                            max_features=50000, sublinear_tf=True)),
  ("clf",   LogisticRegression(max_iter=2000, C=4.0, class_weight="balanced"))])
pipe.fit(train.x, train.y)                      # official train split (41,157)
pred = pipe.predict(test.x)                     # official test split (3,798)
joblib.dump(pipe, "models/sentiment_tfidf_lr.joblib")
json.dump(metrics, open("models/metrics.json", "w"))''', size=8)
    sub(doc, "3.5 Dashboard filters (app/streamlit_app.py)")
    add_code_block(doc, r'''region = st.sidebar.radio("Region", ["Global", "India", "Maharashtra", "Mumbai"])
kw     = st.sidebar.text_input("Keyword / #hashtag search")
f = t[(t.day.dt.date >= d0) & (t.day.dt.date <= d1)]
f = {"Global": f, "India": f[f.is_india], "Maharashtra": f[f.is_maha],
     "Mumbai": f[f.is_mumbai]}[region]
if kw.startswith("#"): f = f[f.tags.str.contains(rf"\b{kw[1:]}\b")]
elif kw:               f = f[f.text.str.lower().str.contains(kw.lower(), regex=False)]
tabs = st.tabs(["Trends", "Sentiment", "Topics", "Keywords & Hashtags",
                "Cases correlation", "Try the model", "Data"])''', size=8)
    body(doc, "Every chart is a Plotly figure, so it supports hover, zoom and legend toggling. The data tab lets the user "
         "download the filtered tweets as CSV. The \u201cTry the model\u201d tab loads the exported joblib pipeline and "
         "classifies any typed text.")

    # TECH
    section_heading(doc, "Technologies Used")
    for n, r in [
        ("Python 3.12, pandas, NumPy, PyArrow", "Data cleaning, joins, aggregation and Parquet storage."),
        ("VADER (vaderSentiment)", "Rule-based sentiment scoring designed for social-media text."),
        ("scikit-learn", "TF-IDF, NMF topic modelling, logistic regression, metrics."),
        ("joblib", "Export of the trained sentiment pipeline (models/sentiment_tfidf_lr.joblib)."),
        ("Streamlit", "Interactive web dashboard with sidebar filters and tabs."),
        ("Plotly", "Interactive bar, line, area, pie, scatter and heat-map charts."),
        ("Playwright (Chromium)", "Headless browser used to capture the dashboard screenshots."),
        ("Kaggle / GitHub / Hugging Face", "Sources of the tweet datasets."),
        ("covid19india.org archive, Our World in Data", "Official COVID-19 case time series."),
    ]:
        bullet(doc, n, r)

    # 4 RESULTS
    section_heading(doc, "4. Results")
    sub(doc, "4.1 Key findings")
    for o in [
        "Sentiment: of 178,484 tweets, 38.6% are Positive, 34.7% Neutral and 26.7% Negative (VADER; mean compound +0.060). "
        "Indian tweets are less negative (20.9% Negative, mean +0.099); Mumbai tweets are mostly Neutral (43.0%), with 20.9% Negative and mean +0.071. "
        "Mean compound fell slightly from +0.066 in July to +0.058 in August.",
        "Health terms: \u201cdeath\u201d (11,873 tweets), \u201ctest\u201d (8,893), \u201cmask\u201d (6,265), \u201cvaccine\u201d (4,086), "
        "\u201chospital\u201d (3,100) and \u201clockdown\u201d (3,072) dominate; symptom words are rare (cough 232, fever 126, oxygen 112), "
        "and flu/influenza appears in 734 tweets. Vaccine mentions rose from 1.66% of tweets in July to 2.55% in August.",
        "India vs rest of world: Indian tweets mention testing (11.3% vs 4.2%), hospitals (4.2% vs 1.4%) and deaths (10.6% vs 6.1%) more often, "
        "but masks less often (1.4% vs 3.8%) and flu/influenza far less (0.11% vs 0.45%).",
        "Hashtags: after #covid19 (99,825) and #coronavirus (10,149), the top tags are #pandemic (1,616), #covid, #india (1,193), #corona, #trump and #lockdown (963). "
        "In India, #odisha (484), #blood (300, plasma/blood-donation appeals) and #indiafightscorona (259) stand out; in Mumbai, #blood (281), #mumbai (68) and #maharashtra (46).",
        "Topics (NMF, k = 8): a broad \u201chealth / vaccine / Trump / news\u201d topic covers 69.1% of tweets; the focused topics are case updates (6.7%), "
        "masks (5.4%), India daily spikes (5.4%), testing (4.6%), spread/symptom-tracking (4.4%) and deaths (2.8%). "
        "For Indian tweets the \u201cIndia daily spike\u201d topic rises to 17.4% and testing to 14.7%. Topic shares were stable between July and August.",
        "Cases vs tweets: India\u2019s daily new cases rose from 50,072 (25 Jul) to 79,461 (30 Aug), Maharashtra from 9,251 to 16,408, while Mumbai stayed "
        "around 700\u20131,300 per day. Tweet volume shows no meaningful relationship with cases: the strongest lag correlations are weak and negative "
        "(India r = \u22120.32 at lag +1 day; Maharashtra r = \u22120.35 at +1; Mumbai r = \u22120.42 at \u22127; same-day r between \u22120.13 and +0.03; n = 25 days).",
    ]:
        dot(doc, o)
    sub(doc, "4.2 Sentiment model")
    body(doc, "The TF-IDF + logistic regression model was evaluated on the dataset\u2019s official test split (3,798 tweets, trained on 41,157): "
         "accuracy 0.782 and macro F1 0.759 (F1: Negative 0.801, Neutral 0.660, Positive 0.815). The confusion matrix "
         "(rows = true Negative, Neutral, Positive) is [[1287, 151, 195], [111, 436, 72], [184, 115, 1247]]. "
         "Applied to the 178,484 COVID19 tweets, it agrees with VADER on 73.4% of tweets and gives a very similar mix "
         "(38.8% Positive, 34.4% Neutral, 26.8% Negative). VADER alone scores 0.935 accuracy on the same test split, "
         "so although the corpus is described as manually tagged, its labels closely follow a lexicon-style rule (see caveats).")
    sub(doc, "4.3 Dashboard screenshots")
    body(doc, "Screenshots of the running dashboard follow, one below another at full width (the filter sidebar is cropped "
         "for readability; it is visible in the live app).")
    for f, c in [
        ("01-trends.png", "Fig. 1: Trends tab (Global) \u2013 daily tweets vs world new cases, weekly tweets by region, and India/Maharashtra/Mumbai share of tweets"),
        ("01-trends-mumbai.png", "Fig. 2: Trends tab with Region = Mumbai \u2013 2,488 tweets vs Mumbai daily new cases"),
        ("02-sentiment.png", "Fig. 3: Sentiment tab \u2013 daily sentiment mix, sentiment by region, 3-day rolling compound score and extreme tweets"),
        ("02-sentiment-mumbai.png", "Fig. 4: Sentiment tab with Region = Mumbai"),
        ("03-topics.png", "Fig. 5: Topics tab \u2013 NMF topic share by week and top terms per topic"),
        ("04-keywords.png", "Fig. 6: Keywords & Hashtags tab \u2013 top hashtags, health-term mentions, terms over time, India vs rest of world"),
        ("05-cases-correlation.png", "Fig. 7: Cases correlation tab \u2013 lag correlation (\u22127 to +7 days), scatter of India tweets vs cases, case context"),
        ("06-try-the-model.png", "Fig. 8: Try the model tab \u2013 live prediction (Negative, 56.9%) and confusion matrix on the test split"),
    ]:
        figure(doc, f, c)
    sub(doc, "4.4 Limitations and caveats")
    for o in [
        "The tweet dataset covers only 25 days (25 Jul \u2013 30 Aug 2020) and its collector capped most days at about 7,500 tweets, with several days missing. "
        "Absolute tweet volume therefore reflects the collection process, not public interest, which explains the weak correlations.",
        "The India share of tweets swings from 0.9% to 18.6% between days, most likely because of the hour at which each batch was collected.",
        "Location is self-reported free text; the regex flags are approximate. \u201cMumbai\u201d includes Thane and Navi Mumbai (Mumbai Metropolitan Region).",
        "The labelled sentiment corpus is from March\u2013April 2020 and VADER matches its labels with 0.935 accuracy, so the model mainly learns to imitate a lexicon; "
        "its scores should be seen as a consistency check rather than ground truth.",
        "Influenza is covered only through flu/influenza keyword mentions (734 tweets); no separate influenza tweet dataset was used.",
        "Correlation is not causation, and 25 data points are too few for a reliable lag estimate.",
    ]:
        dot(doc, o)

    # 5 CONCLUSION
    section_heading(doc, "5. Conclusion")
    body(doc, "This project built a complete, reproducible data-science workflow on real COVID-19 social-media data: "
         "acquisition of 179,108 tweets and official case series for India, Maharashtra, Mumbai and the world; cleaning "
         "and region flagging; sentiment, hashtag, health-term and NMF topic trend analyses; a lag-correlation study; and "
         "a TF-IDF + logistic regression sentiment model (accuracy 0.782, macro F1 0.759) exported with joblib. An "
         "interactive Streamlit dashboard with date, region, keyword and sentiment filters makes every analysis explorable.")
    body(doc, "The findings show that Indian tweets focused more on testing, hospitals and daily case spikes and were less "
         "negative than the rest of the world, that vaccine talk grew during August 2020, and that tweet counts from a "
         "capped collection cannot predict case numbers. Future work includes a longer, uncapped tweet stream (for example "
         "the full 2020\u20132022 period, including India\u2019s 2021 second wave), Hindi and Marathi tweets, transformer-based "
         "sentiment models and a dedicated influenza dataset.")

    # 6 REFERENCES
    section_heading(doc, "6. References")
    refs = [
        "G. Preda, \u201cCOVID19 Tweets,\u201d Kaggle dataset, 2020, https://www.kaggle.com/datasets/gpreda/covid19-tweets; mirror https://github.com/gabrielpreda/covid-19-tweets.",
        "A. Miglani (datatattle), \u201cCoronavirus tweets NLP \u2013 Text Classification,\u201d Kaggle, 2020, https://www.kaggle.com/datasets/datatattle/covid-19-nlp-text-classification; mirror https://huggingface.co/datasets/argilla/twitter-coronavirus.",
        "covid19india.org, COVID-19 India API archive (case_time_series.csv, states.csv, districts.csv), https://data.covid19india.org/.",
        "E. Mathieu et al., \u201cCoronavirus Pandemic (COVID-19),\u201d Our World in Data, https://ourworldindata.org/coronavirus; data https://github.com/owid/covid-19-data.",
        "C. J. Hutto and E. Gilbert, \u201cVADER: A Parsimonious Rule-based Model for Sentiment Analysis of Social Media Text,\u201d Proc. ICWSM, 2014.",
        "D. D. Lee and H. S. Seung, \u201cLearning the parts of objects by non-negative matrix factorization,\u201d Nature, vol. 401, pp. 788\u2013791, 1999.",
        "F. Pedregosa et al., \u201cScikit-learn: Machine Learning in Python,\u201d Journal of Machine Learning Research, vol. 12, pp. 2825\u20132830, 2011.",
        "J. M. Banda et al., \u201cA Large-Scale COVID-19 Twitter Chatter Dataset for Open Scientific Research \u2013 An International Collaboration,\u201d Epidemiologia, vol. 2, no. 3, pp. 315\u2013324, 2021.",
        "E. Chen, K. Lerman and E. Ferrara, \u201cTracking Social Media Discourse About the COVID-19 Pandemic: Development of a Public Coronavirus Twitter Data Set,\u201d JMIR Public Health and Surveillance, vol. 6, no. 2, e19273, 2020.",
        "Streamlit documentation, https://docs.streamlit.io/; Plotly Python documentation, https://plotly.com/python/.",
        "W. McKinney, \u201cData Structures for Statistical Computing in Python,\u201d Proc. 9th Python in Science Conf., 2010 (pandas).",
        "Shree L. R. Tiwari College Of Engineering, Department of Computer Engineering \u2013 Mini Project guidelines (internal).",
    ]
    for i, r in enumerate(refs, 1):
        add_para(doc, f"{i}. {r}", size=12, align="justify", space_after=6, line_spacing=1.3)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print("Wrote", OUT)


if __name__ == "__main__":
    build()
